from datetime import datetime
from io import BytesIO
from zipfile import BadZipFile

import pytest
from openpyxl import Workbook

from src.infra.excel_readers.openpyxl_invoice_reader import OpenpyxlInvoiceReader


@pytest.fixture
def invoice_workbook() -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(
        [
            'Категория',
            'Крайний момент оплаты',
            'Дата счета',
            'Предоплата до загрузки',
            'Юр лицо',
            'Счет на',
            'Полный номер счета',
            'Сервис для срока оплаты',
            'Рынок',
            'Сотрудник',
        ]
    )
    worksheet.append(['Предоплата', 'До загрузки', datetime(2026, 10, 20), 'да'])
    worksheet.append(['Доплата', 'После загрузки', datetime(2026, 10, 21), 'да'])
    return workbook


@pytest.fixture
def excel_bytes(invoice_workbook: Workbook) -> bytes:
    with BytesIO() as excel_buffer:
        invoice_workbook.save(excel_buffer)
        return excel_buffer.getvalue()


@pytest.fixture
def invoice_reader() -> OpenpyxlInvoiceReader:
    return OpenpyxlInvoiceReader()


def test_load_exposes_workbook(invoice_reader, excel_bytes):
    invoice_reader.load(excel_bytes)
    assert isinstance(invoice_reader.workbook, Workbook)


def test_load_rejects_empty_bytes(invoice_reader):
    with pytest.raises(BadZipFile):
        invoice_reader.load(b'')


def test_unload_clears_workbook(invoice_reader, excel_bytes):
    invoice_reader.load(excel_bytes)
    invoice_reader.unload()
    with pytest.raises(ValueError, match='Файл не был загружен'):
        _ = invoice_reader.workbook


def test_reload_creates_new_workbook(invoice_reader, excel_bytes):
    invoice_reader.load(excel_bytes)
    previous_workbook = invoice_reader.workbook
    invoice_reader.unload()
    invoice_reader.load(excel_bytes)
    assert invoice_reader.workbook is not previous_workbook


def test_load_rejects_already_loaded_workbook(invoice_reader, excel_bytes):
    invoice_reader.load(excel_bytes)
    with pytest.raises(RuntimeError, match='Файл уже был загружен'):
        invoice_reader.load(excel_bytes)


def test_read_invoices_returns_invoices_in_row_order(
    invoice_reader, excel_bytes, invoice_factory
):
    invoice_reader.load(excel_bytes)

    assert invoice_reader.read_invoices(worksheet_index=0) == [
        invoice_factory(
            category='Предоплата',
            payment_terms='До загрузки',
            invoice_date=datetime(2026, 10, 20),
            sign='да',
        ),
        invoice_factory(
            category='Доплата',
            payment_terms='После загрузки',
            invoice_date=datetime(2026, 10, 21),
            sign='да',
        ),
    ]


def test_read_invoices_uses_selected_worksheet(
    invoice_reader, invoice_workbook, invoice_factory
):
    worksheet = invoice_workbook.create_sheet('Другие счета')
    worksheet.append(
        [
            'Дата счета',
            'Предоплата до загрузки',
            'Примечание',
            'Крайний момент оплаты',
            'Категория',
            'Юр лицо',
            'Счет на',
            'Полный номер счета',
            'Сервис для срока оплаты',
            'Рынок',
            'Сотрудник',
        ]
    )
    worksheet.append(
        [datetime(2026, 10, 22), 'да', 'Срочно', 'До загрузки', 'Предоплата']
    )
    with BytesIO() as excel_buffer:
        invoice_workbook.save(excel_buffer)
        invoice_reader.load(excel_buffer.getvalue())

    assert invoice_reader.read_invoices(worksheet_index=1) == [
        invoice_factory(
            category='Предоплата',
            payment_terms='До загрузки',
            invoice_date=datetime(2026, 10, 22),
            sign='да',
        ),
    ]


def test_read_invoices_returns_empty_list_for_headers_only(
    invoice_reader, invoice_workbook
):
    invoice_workbook.active.delete_rows(2, 2)
    with BytesIO() as excel_buffer:
        invoice_workbook.save(excel_buffer)
        invoice_reader.load(excel_buffer.getvalue())

    assert invoice_reader.read_invoices(worksheet_index=0) == []


@pytest.mark.parametrize(
    ('missing_column', 'missing_field'),
    [
        (1, 'category'),
        (2, 'payment_terms'),
        (3, 'invoice_date'),
        (4, 'sign'),
        (5, 'legal_entity'),
        (6, 'subject'),
        (7, 'invoice_number'),
    ],
)
def test_read_invoices_preserves_missing_columns_as_none(
    invoice_reader,
    invoice_workbook,
    missing_column,
    missing_field,
):
    invoice_workbook.active.delete_cols(missing_column)
    with BytesIO() as excel_buffer:
        invoice_workbook.save(excel_buffer)
        invoice_reader.load(excel_buffer.getvalue())

    invoices = invoice_reader.read_invoices(worksheet_index=0)
    assert len(invoices) == 2
    assert all(getattr(invoice, missing_field) is None for invoice in invoices)


@pytest.mark.parametrize(
    'empty_columns',
    [(column,) for column in range(1, 8)] + [tuple(range(1, 11))],
)
def test_read_invoices_preserves_missing_values(
    invoice_reader, invoice_workbook, empty_columns, invoice_factory
):
    worksheet = invoice_workbook.active
    for column_index in empty_columns:
        worksheet.cell(row=2, column=column_index).value = None
    expected_values = [
        worksheet.cell(row=2, column=column_index).value
        for column_index in range(1, 11)
    ]
    with BytesIO() as excel_buffer:
        invoice_workbook.save(excel_buffer)
        invoice_reader.load(excel_buffer.getvalue())

    invoices = invoice_reader.read_invoices(worksheet_index=0)

    assert len(invoices) == 2
    assert invoices[0] == invoice_factory(*expected_values)
    assert invoices[1] == invoice_factory(
        'Доплата', 'После загрузки', datetime(2026, 10, 21), 'да'
    )


def test_read_invoices_preserves_blank_rows_between_invoices(
    invoice_reader, invoice_workbook, invoice_factory
):
    invoice_workbook.active.insert_rows(3)
    with BytesIO() as excel_buffer:
        invoice_workbook.save(excel_buffer)
        invoice_reader.load(excel_buffer.getvalue())

    assert invoice_reader.read_invoices(worksheet_index=0) == [
        invoice_factory('Предоплата', 'До загрузки', datetime(2026, 10, 20), 'да'),
        invoice_factory(None, None, None, None),
        invoice_factory('Доплата', 'После загрузки', datetime(2026, 10, 21), 'да'),
    ]


@pytest.mark.parametrize('sign', ['да', 'Да', 'нет', None])
def test_read_invoices_preserves_sign_without_filtering(
    invoice_reader, invoice_workbook, sign
):
    invoice_workbook.active.cell(row=2, column=4).value = sign
    with BytesIO() as excel_buffer:
        invoice_workbook.save(excel_buffer)
        invoice_reader.load(excel_buffer.getvalue())

    invoices = invoice_reader.read_invoices(worksheet_index=0)

    assert len(invoices) == 2
    assert invoices[0].sign == sign


def test_read_invoices_maps_invoice_details(
    invoice_reader, invoice_workbook, invoice_factory
):
    worksheet = invoice_workbook.active
    worksheet.cell(row=2, column=5).value = 'ООО Перевозчик'
    worksheet.cell(row=2, column=6).value = 'Перевозка груза'
    worksheet.cell(row=2, column=7).value = 'INV-001'
    worksheet.cell(row=2, column=8).value = 'Отдел перевозок'
    worksheet.cell(row=2, column=9).value = 'РБ'
    worksheet.cell(row=2, column=10).value = 'employee@example.com'
    with BytesIO() as excel_buffer:
        invoice_workbook.save(excel_buffer)
        invoice_reader.load(excel_buffer.getvalue())

    assert invoice_reader.read_invoices(worksheet_index=0)[0] == invoice_factory(
        category='Предоплата',
        payment_terms='До загрузки',
        invoice_date=datetime(2026, 10, 20),
        sign='да',
        legal_entity='ООО Перевозчик',
        subject='Перевозка груза',
        invoice_number='INV-001',
        payment_service='Отдел перевозок',
        market='РБ',
        employee='employee@example.com',
    )


def test_read_invoices_requires_loaded_workbook(invoice_reader):
    with pytest.raises(ValueError, match='Файл не был загружен'):
        invoice_reader.read_invoices(worksheet_index=0)


def test_read_invoices_rejects_missing_worksheet(invoice_reader, excel_bytes):
    invoice_reader.load(excel_bytes)
    with pytest.raises(IndexError):
        invoice_reader.read_invoices(worksheet_index=1)
