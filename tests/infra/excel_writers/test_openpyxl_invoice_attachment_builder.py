from dataclasses import fields
from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

from src.domain.entity.invoice import Invoice
from src.infra.excel_writers.openpyxl_invoice_attachment_builder import (
    OpenpyxlInvoiceAttachmentBuilder,
)


def test_build_attachment_writes_invoice_fields_and_metadata(invoice_factory):
    invoice = invoice_factory(
        client_id='CLIENT-001',
        invoice_number='INV-001',
        invoice_amount=125.5,
        invoice_date=datetime(2026, 10, 20),
        currency='EUR',
        market='РБ',
        employee='employee@example.com',
        payment_service='Авиа',
        sign='да',
    )

    attachment = OpenpyxlInvoiceAttachmentBuilder().build_attachment([invoice])

    assert attachment.filename == 'invoices.xlsx'
    assert (
        attachment.content_type
        == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    workbook = load_workbook(BytesIO(attachment.content), data_only=True)
    try:
        worksheet = workbook['Счета']
        rows = list(worksheet.iter_rows(values_only=True))
        assert len(rows) == 2
        assert len(rows[0]) == len(fields(Invoice))
        values = dict(zip(rows[0], rows[1]))
        assert values['ID клиента'] == 'CLIENT-001'
        assert values['Полный номер счета'] == 'INV-001'
        assert values['Сумма счета'] == 125.5
        assert values['Дата счета'] == datetime(2026, 10, 20)
        assert values['Валюта'] == 'EUR'
        assert values['Рынок'] == 'РБ'
        assert values['Сотрудник'] == 'employee@example.com'
        assert values['Сервис для срока оплаты'] == 'Авиа'
        assert values['Предоплата до загрузки'] == 'да'
        assert values['Юр лицо'] is None
    finally:
        workbook.close()


def test_build_attachment_keeps_consecutive_workbooks_independent(invoice_factory):
    builder = OpenpyxlInvoiceAttachmentBuilder()
    first_attachment = builder.build_attachment(
        [invoice_factory(invoice_number='INV-001')]
    )
    second_attachment = builder.build_attachment(
        [invoice_factory(invoice_number='INV-002')]
    )

    for attachment, expected_number in [
        (first_attachment, 'INV-001'),
        (second_attachment, 'INV-002'),
    ]:
        workbook = load_workbook(BytesIO(attachment.content), data_only=True)
        try:
            rows = list(workbook.active.iter_rows(values_only=True))
            assert len(rows) == 2
            assert dict(zip(rows[0], rows[1]))['Полный номер счета'] == expected_number
        finally:
            workbook.close()


def test_build_attachment_supports_empty_invoice_list():
    attachment = OpenpyxlInvoiceAttachmentBuilder().build_attachment([])
    workbook = load_workbook(BytesIO(attachment.content), data_only=True)
    try:
        assert workbook.active.max_row == 1
        assert workbook.active.max_column == len(fields(Invoice))
    finally:
        workbook.close()
