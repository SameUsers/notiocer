from datetime import datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from src.application.commands.create_mailing_list import CreateMailingListCommand
from src.application.interactors.create_mailing_list import CreateMailingListInteractor
from src.domain.services.invoice_message_builder import InvoiceMessageBuilder
from src.domain.services.mailing_invoice_selector import MailingInvoiceSelector
from src.domain.services.recipient_mailing_list_builder import RecipientMailingListBuilder
from src.infra.excel_readers.openpyxl_invoice_reader import OpenpyxlInvoiceReader
from src.infra.excel_writers.openpyxl_invoice_attachment_builder import OpenpyxlInvoiceAttachmentBuilder


@pytest.fixture
def invoice_report():
    workbook = Workbook()
    try:
        worksheet = workbook.active
        worksheet.append([
            "Категория", "Дата счета", "Признак", "Сервис для срока оплаты",
            "Рынок", "Сотрудник", "Полный номер счета",
        ])
        worksheet.append(["Д0", datetime(2000, 1, 1), "да", "Авиа", "РБ", "Сотрудник", "INV-001"])
        worksheet.append(["Д0", datetime(2000, 1, 1), "да", "Авиа", "РФ", "Сотрудник", "INV-002"])
        with BytesIO() as buffer:
            workbook.save(buffer)
            return buffer.getvalue()
    finally:
        workbook.close()


def create_interactor(reader, attachment_builder):
    return CreateMailingListInteractor(
        invoice_reader=reader,
        invoice_selector=MailingInvoiceSelector({"Д0"}, timedelta(days=3)),
        mailing_list_builder=RecipientMailingListBuilder(),
        message_builder=InvoiceMessageBuilder(),
        attachment_builder=attachment_builder,
    )


def test_create_mailing_list_attaches_only_each_recipients_invoices(invoice_report):
    reader = OpenpyxlInvoiceReader()
    interactor = create_interactor(reader, OpenpyxlInvoiceAttachmentBuilder())

    messages = interactor.execute(CreateMailingListCommand(invoice_report))

    assert {message.recipient_email for message in messages} == {
        "CM_Air_BY@stalogistic.com", "cm-air-ru@stalogistic.com",
    }
    expected_numbers = {"CM_Air_BY@stalogistic.com": "INV-001", "cm-air-ru@stalogistic.com": "INV-002"}
    for message in messages:
        workbook = load_workbook(BytesIO(message.attachment.content), data_only=True)
        try:
            rows = list(workbook.active.iter_rows(values_only=True))
            assert len(rows) == 2
            assert dict(zip(rows[0], rows[1]))["Полный номер счета"] == expected_numbers[message.recipient_email]
        finally:
            workbook.close()
    with pytest.raises(ValueError):
        _ = reader.workbook



def test_attachment_failure_propagates_and_unloads_reader(invoice_report, mocker):
    reader = OpenpyxlInvoiceReader()
    attachment_builder = OpenpyxlInvoiceAttachmentBuilder()
    attachment_error = OSError("Cannot create attachment")
    mocker.patch.object(attachment_builder, "build_attachment", side_effect=attachment_error)
    interactor = create_interactor(reader, attachment_builder)

    with pytest.raises(OSError) as error:
        interactor.execute(CreateMailingListCommand(invoice_report))

    assert error.value is attachment_error
    with pytest.raises(ValueError):
        _ = reader.workbook
