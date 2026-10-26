from collections.abc import Sequence
from dataclasses import Field, fields
from io import BytesIO
from typing import Annotated, Any, get_args, get_origin

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from notiocer.application.dto.email_attachment import EmailAttachment
from notiocer.domain.models.invoice import Invoice


class OpenpyxlInvoiceAttachmentBuilder:
    """Создаёт XLSX-вложение со всеми полями счетов и оформленными заголовками."""

    def _style_headers(self, worksheet: Worksheet) -> None:
        """Оформляет заголовки и задаёт размеры строк и колонок."""
        header_fill = PatternFill(fill_type='solid', fgColor='1F4E78')
        header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        header_alignment = Alignment(
            horizontal='center', vertical='center', wrap_text=True
        )
        border_side = Side(style='thin', color='D9E2F3')

        for column_index, cell in enumerate(worksheet[1], start=1):
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = Border(
                left=border_side,
                right=border_side,
                top=border_side,
                bottom=border_side,
            )
            column_letter = get_column_letter(column_index)
            worksheet.column_dimensions[column_letter].width = 22

        worksheet.row_dimensions[1].height = 40

    @staticmethod
    def _column_title(invoice_field: Field[Any]) -> str:
        """Берёт заголовок из Annotated или использует имя поля."""
        annotation = invoice_field.type
        if get_origin(annotation) is Annotated:
            annotation_arguments = get_args(annotation)
            if len(annotation_arguments) > 1 and isinstance(
                annotation_arguments[1], str
            ):
                return annotation_arguments[1]
        return invoice_field.name

    def build(self, invoices: Sequence[Invoice]) -> EmailAttachment:
        """Создаёт XLSX в памяти и освобождает книгу после сохранения."""
        invoice_fields = fields(Invoice)
        workbook = Workbook()
        try:
            worksheet = workbook.active
            if worksheet is None:
                raise RuntimeError('Не удалось создать лист для вложения')
            worksheet.title = 'Счета'

            worksheet.append(
                [self._column_title(invoice_field) for invoice_field in invoice_fields]
            )
            self._style_headers(worksheet)

            field_names = [invoice_field.name for invoice_field in invoice_fields]
            for row_index, invoice in enumerate(invoices, start=2):
                row_values = [
                    getattr(invoice, field_name) for field_name in field_names
                ]
                worksheet.append(row_values)
                for column_index, value in enumerate(row_values, start=1):
                    if isinstance(value, str):
                        worksheet.cell(row_index, column_index).data_type = 's'

            with BytesIO() as attachment_buffer:
                workbook.save(attachment_buffer)
                attachment_content = attachment_buffer.getvalue()
        finally:
            workbook.close()
        return EmailAttachment(
            filename='invoices.xlsx',
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            content=attachment_content,
        )
