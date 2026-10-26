from collections.abc import Sequence
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook

from notiocer.domain.models.invoice import Invoice
from dataclasses import fields
from typing import Any, get_args, get_type_hints


class OpenpyxlInvoiceReader:
    """Читает счета из XLSX через openpyxl с циклом load → read_invoices → unload."""

    def __init__(self) -> None:
        """Создаёт reader без загруженной книги."""
        self._workbook: Workbook | None = None

    def load(self, report_content: bytes) -> None:
        """Загружает книгу из байтов Excel-файла.

        Raises:
            RuntimeError: Предыдущая книга ещё не освобождена.
        """
        if self._workbook is not None:
            raise RuntimeError('Файл уже был загружен. Сначала вызовите unload()')
        self._workbook = load_workbook(filename=BytesIO(report_content), data_only=True)

    def _map_row_to_invoice(
        self,
        values_by_header: dict[str, Any],
    ) -> Invoice:
        """Создаёт Invoice, используя метаданные Annotated."""

        hints = get_type_hints(Invoice, include_extras=True)

        return Invoice(
            **{
                field.name: values_by_header.get(get_args(hints[field.name])[1])
                for field in fields(Invoice)
            }
        )

    def _read_headers(self, worksheet_index: int) -> Sequence[Any]:
        """Читает первую строку листа как заголовки."""
        return self._read_row(row_index=0, worksheet_index=worksheet_index)

    def _read_row(self, row_index: int, worksheet_index: int) -> Sequence[Any]:
        """Читает строку по индексам строки и листа, начиная с нуля."""
        worksheet = self.workbook.worksheets[worksheet_index]
        return [cell.value for cell in worksheet[row_index + 1]]

    def _get_row_count(self, worksheet_index: int) -> int:
        """Возвращает количество строк на листе."""
        worksheet = self.workbook.worksheets[worksheet_index]
        return worksheet.max_row

    def read_invoices(self, worksheet_index: int) -> list[Invoice]:
        """Преобразует строки после заголовка в счета без бизнес-фильтрации."""
        row_count = self._get_row_count(worksheet_index=worksheet_index)
        headers = self._read_headers(worksheet_index=worksheet_index)
        rows_by_header = [
            dict(
                zip(
                    headers,
                    self._read_row(
                        row_index=row_index, worksheet_index=worksheet_index
                    ),
                )
            )
            for row_index in range(1, row_count)
        ]
        return [
            self._map_row_to_invoice(values_by_header)
            for values_by_header in rows_by_header
        ]

    def unload(self) -> None:
        """Освобождает загруженную книгу и сбрасывает состояние reader."""
        if self._workbook is not None:
            self._workbook.close()
            self._workbook = None

    @property
    def workbook(self) -> Workbook:
        """Возвращает загруженную книгу.

        Raises:
            ValueError: Книга ещё не загружена.
        """
        if self._workbook is None:
            raise ValueError('Файл не был загружен')
        return self._workbook
