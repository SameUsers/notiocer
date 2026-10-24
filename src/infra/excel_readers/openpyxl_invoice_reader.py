from collections.abc import Sequence
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook

from src.domain.entity.invoice import Invoice


class OpenpyxlInvoiceReader:
    """Читает счета из книги Excel с помощью openpyxl.

    Ридер следует жизненному циклу загрузка/чтение/выгрузка: вызовите
    :meth:`load` с сырыми байтами книги, прочитайте счета с помощью
    :meth:`read_invoices` и освободите книгу с помощью :meth:`unload`.
    """

    def __init__(self) -> None:
        """Инициализирует ридер без загруженной книги."""
        self._workbook: Workbook | None = None

    def load(self, excel_bytes: bytes) -> None:
        """Загружает книгу из сырых байтов Excel в память.

        Args:
            excel_bytes: Сырое содержимое файла ``.xlsx`` для разбора.

        Raises:
            RuntimeError: Если книга уже загружена. Сначала вызовите
                :meth:`unload`.
        """
        if self._workbook is not None:
            raise RuntimeError('Файл уже был загружен. Сначала вызовите unload()')
        self._workbook = load_workbook(filename=BytesIO(excel_bytes), data_only=True)

    def _map_row_to_invoice(self, row_values: dict[str, Any]) -> Invoice:
        """Преобразует строку с ключами-заголовками в счёт.

        Args:
            row_values: Отображение заголовка колонки на значение ячейки для строки.

        Returns:
            Счёт, построенный из значений строки.
        """
        return Invoice(
            client_id=row_values.get('ID клиента'),
            legal_entity=row_values.get('Юр лицо'),
            subject=row_values.get('Счет на'),
            invoice_number=row_values.get('Полный номер счета'),
            invoice_date=row_values.get('Дата счета'),
            invoice_amount=row_values.get('Сумма счета'),
            balance=row_values.get('Сальдо'),
            currency=row_values.get('Валюта'),
            category=row_values.get('Категория'),
            payment_due_date=row_values.get('Срок оплаты'),
            sign=row_values.get('Предоплата до загрузки'),
            employee=row_values.get('Сотрудник'),
            equipment_service=row_values.get('Услуга отдела оборудования'),
            transport_reference=row_values.get('Номер контейнера / AWB /авто'),
            unloading_date=row_values.get('Дата выгрузки'),
            sfb_curator=row_values.get('Куратор СФБ'),
            shipping_date=row_values.get('Дата отгрузки'),
            estimated_arrival_date=row_values.get('Прогнозная ETA'),
            market=row_values.get('Рынок'),
            payment_service=row_values.get('Сервис для срока оплаты'),
            payment_terms_type=row_values.get('Тип платежных условий'),
            contract_details=row_values.get('№ и дата договора'),
            service_type=row_values.get('Тип сервиса'),
            partial_prepayment=row_values.get('Частичная предоплата'),
            partial_prepayment_terms=row_values.get('Условия частичной предоплаты'),
            payment_terms=row_values.get('Крайний момент оплаты'),
            calculated_payment_term=row_values.get('Расчетный с/о'),
        )

    def _read_headers(self, worksheet_index: int) -> Sequence[Any]:
        """Читает строку заголовков листа.

        Args:
            worksheet_index: Индекс листа (начиная с нуля).

        Returns:
            Значения ячеек строки заголовков.
        """
        return self._read_row(row_index=0, worksheet_index=worksheet_index)

    def _read_row(self, row_index: int, worksheet_index: int) -> Sequence[Any]:
        """Читает одну строку с листа.

        Args:
            row_index: Индекс строки (начиная с нуля) в пределах диапазона данных.
            worksheet_index: Индекс листа (начиная с нуля).

        Returns:
            Значения ячеек указанной строки.
        """
        worksheet = self.workbook.worksheets[worksheet_index]
        return [cell.value for cell in worksheet[row_index + 1]]

    def _get_row_count(self, worksheet_index: int) -> int:
        """Возвращает количество строк на листе.

        Args:
            worksheet_index: Индекс листа (начиная с нуля).

        Returns:
            Максимальное количество строк на листе.
        """
        worksheet = self.workbook.worksheets[worksheet_index]
        return worksheet.max_row

    def read_invoices(self, worksheet_index: int) -> list[Invoice]:
        """Читает все счета с листа.

        Пропускает строку заголовков и преобразует каждую последующую
        строку в счёт.

        Args:
            worksheet_index: Индекс листа (начиная с нуля) для чтения.

        Returns:
            Список счетов, разобранных с листа.
        """
        row_count = self._get_row_count(worksheet_index=worksheet_index)
        headers = self._read_headers(worksheet_index=worksheet_index)
        invoice_rows = [
            dict(
                zip(
                    headers,
                    self._read_row(
                        row_index=row_index, worksheet_index=worksheet_index
                    ),
                )
            )
            for row_index in range(1, row_count)
        ]
        return [self._map_row_to_invoice(row_values) for row_values in invoice_rows]

    def unload(self) -> None:
        """Закрывает загруженную книгу и сбрасывает состояние ридера."""
        if self._workbook is not None:
            self._workbook.close()
            self._workbook = None

    @property
    def workbook(self) -> Workbook:
        """Текущая загруженная книга.

        Returns:
            Загруженная книга openpyxl.

        Raises:
            ValueError: Если книга не была загружена.
        """
        if self._workbook is None:
            raise ValueError('Файл не был загружен')
        return self._workbook
