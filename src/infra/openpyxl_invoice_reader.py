from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook

from src.domain.entity.invoice import Invoice


class OpenpyxlInvoiceReader:
    def __init__(self) -> None:
        self._workbook: Workbook | None = None

    def load(self, excel_bytes: bytes) -> None:
        if self._workbook is not None:
            raise RuntimeError("Файл уже был загружен. Сначала вызовите unload()")
        self._workbook = load_workbook(filename=BytesIO(excel_bytes), data_only=True)

    def _map_row_to_invoice(self, row_values: dict[str, Any]) -> Invoice:
        return Invoice(
            category=row_values["Категория"],
            payment_terms=row_values["Крайний момент оплаты"],
            invoice_date=row_values["Дата счета"],
            sign=row_values["Признак"],
            invoice_number=row_values["Полный номер счета"],
            subject=row_values["Счет на"],
            legal_entity=row_values["Юр лицо"],
        )

    def _read_headers(self, worksheet_index: int) -> list[Any]:
        return self._read_row(row_index=0, worksheet_index=worksheet_index)

    def _read_row(self, row_index: int, worksheet_index: int) -> list[Any]:
        worksheet = self.workbook.worksheets[worksheet_index]
        return [cell.value for cell in worksheet[row_index + 1]]

    def _get_row_count(self, worksheet_index: int) -> int:
        worksheet = self.workbook.worksheets[worksheet_index]
        return worksheet.max_row

    def read_invoices(self, worksheet_index: int) -> list[Invoice]:
        row_count = self._get_row_count(worksheet_index=worksheet_index)
        headers = self._read_headers(worksheet_index=worksheet_index)
        invoice_rows = [
            dict(zip(headers, self._read_row(row_index=row_index, worksheet_index=worksheet_index)))
            for row_index in range(1, row_count)
        ]
        return [self._map_row_to_invoice(row_values) for row_values in invoice_rows]

    def unload(self) -> None:
        if self._workbook is not None:
            self._workbook.close()
            self._workbook = None

    @property
    def workbook(self) -> Workbook:
        if self._workbook is None:
            raise ValueError("Файл не был загружен")
        return self._workbook
